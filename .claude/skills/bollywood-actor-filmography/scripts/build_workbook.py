#!/usr/bin/env python3
"""Build the Bollywood actor filmography workbook from JSON research files.

The research lives in one directory per actor (one JSON file per sheet); this
script validates it and renders the formatted seven-sheet workbook. Rebuild the
workbook from the JSON after every change instead of editing the .xlsx.

Commands:
  init DATA_DIR --actor NAME [--cutoff YYYY-MM-DD]   create empty research files
  schema [SHEET]                                     list JSON keys and allowed values
  validate DATA_DIR                                  check data without building
  build DATA_DIR -o OUTPUT.xlsx                      validate, then write the workbook

Run with the interpreter that has openpyxl:
  ~/.venvs/filmography/bin/python build_workbook.py ...
"""
from __future__ import annotations

import argparse
import datetime as dt
import difflib
import json
import math
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field as dc_field
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

try:
    import openpyxl
    from openpyxl.chart import BarChart, Reference, ScatterChart, Series
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("openpyxl is not installed for this interpreter. "
             "Run the script with ~/.venvs/filmography/bin/python")

# ---------------------------------------------------------------------------
# Controlled vocabularies (from SKILL.md)
# ---------------------------------------------------------------------------

ROLE_TYPES = ["Lead", "Co-lead", "Ensemble lead", "Supporting", "Antagonist",
              "Extended special appearance", "Cameo", "Self appearance",
              "Voice role", "Narrator", "Child role"]
ROMANCE_PRESENCE = ["None", "Implied", "Minor", "Secondary", "Major", "Central"]
RELATIONSHIP_CLASSES = [
    "Frictionless idealized romance", "United couple versus external world",
    "Friendship to love", "Rivals to lovers", "Playful conflict masking attraction",
    "Genuine incompatibility gradually resolved", "Established relationship",
    "Marriage under strain", "Relationship damaged by mistrust",
    "Relationship manipulated by outsiders", "Forbidden love", "Unrequited love",
    "Love triangle", "Tragic lovers", "Unequal or coercive relationship",
    "Obsessive pursuit", "Romance subordinate to another genre",
    "Multiple romantic relationships", "No meaningful romance"]
YES_NO = ["Yes", "No", "Partial", "Unclear", "Not applicable"]
OPPOSITION_SOURCES = ["Family", "Class", "Caste", "Religion", "Community",
                      "Existing relationship", "Villain", "Crime", "War or politics",
                      "Illness", "Distance", "Fate", "Other"]
ENDINGS = ["Happy", "Tragic", "Unresolved"]
CONFIDENCE = ["High", "Medium", "Low", "Unknown"]
RELIABILITY = ["A", "B", "C", "D"]
CAST_FUNCTIONS = ["Co-lead", "Supporting", "Antagonist", "Mentor", "Family", "Friend",
                  "Comic support", "Romantic rival", "Authority figure", "Cameo",
                  "Self appearance"]
IMPORTANCE = ["Major", "Medium", "Minor"]
WIN_NOM = ["Win", "Nomination"]
COMP_HON = ["Competitive", "Honorary"]

# Overview groupings; every Role Type and Romance Presence value lands in exactly one.
ROLE_GROUPS = [
    ("Lead and co-lead films", ["Lead", "Co-lead", "Ensemble lead"]),
    ("Supporting roles", ["Supporting", "Antagonist"]),
    ("Cameos and special appearances", ["Cameo", "Extended special appearance", "Self appearance"]),
    ("Voice, narration and child roles", ["Voice role", "Narrator", "Child role"]),
]
ROMANCE_GROUPS = [
    ("Films with central romance", ["Central"]),
    ("Films with major romance", ["Major"]),
    ("Films with secondary romance", ["Secondary"]),
    ("Films with minor or implied romance", ["Minor", "Implied"]),
    ("Films without meaningful romance", ["None"]),
]

# ---------------------------------------------------------------------------
# Schema
# ---------------------------------------------------------------------------
# kind: id | ref | refs | text | long | score | date | year | int | num | money
#       | enum | multi | url
# ref/refs name the sheet whose IDs the value must resolve to. Multi-valued
# cells (refs, multi, and name lists) use "; " as the separator; JSON arrays
# are also accepted and joined.

DEFAULT_WIDTH = {"id": 12, "ref": 12, "refs": 22, "text": 20, "long": 50, "score": 10,
                 "date": 12, "year": 9, "int": 10, "num": 11, "money": 15,
                 "enum": 20, "multi": 24, "url": 42}


@dataclass
class Field:
    key: str
    header: str
    kind: str = "text"
    width: int | None = None
    choices: list[str] | None = None
    ref: str | None = None
    analysis: bool = False  # interpretive judgement rather than fact

    @property
    def col_width(self) -> int:
        return self.width or DEFAULT_WIDTH[self.kind]

    @property
    def category(self) -> str:
        if self.kind == "score":
            return "score"
        return "analysis" if self.analysis else "fact"


@dataclass
class SheetSpec:
    key: str
    title: str
    file: str
    prefix: str
    freeze: str
    fields: list[Field]
    recommended: list[str] = dc_field(default_factory=list)  # warn when blank

    def field(self, key: str) -> Field:
        return next(f for f in self.fields if f.key == key)

    def col(self, key: str) -> str:
        return get_column_letter(next(i for i, f in enumerate(self.fields, 1) if f.key == key))


F = Field

FILM_FIELDS = [
    F("film_id", "Film ID", "id"),
    F("title", "Title", width=26),
    F("release_date", "Release Date", "date"),
    F("release_year", "Release Year", "year"),
    F("decade", "Decade", width=9),
    F("language", "Language", width=14),
    F("genre", "Genre", width=16),
    F("secondary_genre", "Secondary Genre", width=16),
    F("runtime", "Runtime", "int"),
    F("director", "Director", width=22),
    F("producers", "Producer or Producers", width=26),
    F("production_banner", "Production Banner", width=24),
    F("writers", "Writer or Writers", width=26),
    F("source_material", "Source Material or Remake Status", width=26),
    F("song_composer", "Song Composer", width=22),
    F("background_score_composer", "Background-Score Composer", width=22),
    F("lyricist", "Lyricist", width=22),
    F("character", "Actor’s Character", width=20),
    F("alternate_identity", "Alternate Identity or Double Role", width=22),
    F("role_type", "Role Type", "enum", choices=ROLE_TYPES),
    F("billing_position", "Billing Position", "num"),
    F("character_occupation", "Character Occupation", width=22),
    F("character_social_background", "Character Social Background", width=26),
    F("role_gist", "Role Gist", "long", width=60, analysis=True),
    F("principal_motivation", "Principal Motivation", "long", width=36, analysis=True),
    F("principal_conflict", "Principal Conflict", "long", width=36, analysis=True),
    F("character_arc", "Character Arc", "long", width=40, analysis=True),
    F("moral_alignment", "Moral Alignment", width=18, analysis=True),
    F("character_outcome", "Character Outcome", "long", width=30),
    F("female_leads", "Principal Female Lead or Leads", width=26),
    F("romance_presence", "Romance Presence", "enum", width=14, choices=ROMANCE_PRESENCE),
    F("romance_centrality", "Romance Centrality Score", "score"),
    F("relationship_classification", "Overall Relationship Classification", "enum",
      width=30, choices=RELATIONSHIP_CLASSES, analysis=True),
    F("budget", "Budget", "money"),
    F("domestic_nett", "Domestic Nett", "money"),
    F("domestic_gross", "Domestic Gross", "money"),
    F("overseas_gross", "Overseas Gross", "money"),
    F("worldwide_gross", "Worldwide Gross", "money"),
    F("currency", "Currency", width=13),
    F("verdict", "Box-Office Verdict", width=16),
    F("annual_rank", "Annual Box-Office Rank", "num"),
    F("critical_reception", "Critical Reception Summary", "long", width=45, analysis=True),
    F("recognition", "Recognition or Legacy", "long", width=40, analysis=True),
    F("source_ids", "Source IDs", "refs", ref="sources"),
    F("confidence", "Confidence Grade", "enum", width=12, choices=CONFIDENCE),
    F("research_notes", "Research Notes", "long", width=45),
]
MONEY_KEYS = ["budget", "domestic_nett", "domestic_gross", "overseas_gross", "worldwide_gross"]

ROMANCE_FIELDS = [
    F("pairing_id", "Pairing ID", "id"),
    F("film_id", "Film ID", "ref", ref="films"),
    F("film", "Film", width=24),
    F("release_year", "Release Year", "year"),
    F("character", "Actor’s Character", width=20),
    F("actress", "Actress", width=20),
    F("actress_character", "Actress’s Character", width=20),
    F("actress_prominence", "Actress Role Prominence", width=18),
    F("principal_female_lead", "Principal Female Lead", "enum", width=14, choices=YES_NO),
    F("relationship_category", "Relationship Category", "enum", width=30,
      choices=RELATIONSHIP_CLASSES, analysis=True),
    F("initial_relationship", "Initial Relationship", width=22),
    F("how_they_meet", "How They Meet", "long", width=36),
    F("who_initiates", "Who Initiates", width=18),
    F("courtship_pattern", "Courtship Pattern", "long", width=36, analysis=True),
    F("emotional_dynamic", "Emotional Dynamic", "long", width=36, analysis=True),
    F("pair_conflict", "Pair’s Principal Conflict", "long", width=36, analysis=True),
    F("pair_friction", "Pair Friction Score", "score"),
    F("pair_friction_explanation", "Pair Friction Explanation", "long", width=40, analysis=True),
    F("external_opposition", "External Opposition Score", "score"),
    F("external_opposition_source", "External Opposition Source", "multi",
      choices=OPPOSITION_SOURCES),
    F("external_opposition_explanation", "External Opposition Explanation", "long",
      width=40, analysis=True),
    F("lover_boy", "Lover-Boy Score", "score"),
    F("lover_boy_explanation", "Lover-Boy Explanation", "long", width=40, analysis=True),
    F("romance_centrality", "Romance Centrality Score", "score"),
    F("relationship_health", "Relationship Health Score", "score"),
    F("mutuality", "Mutuality Score", "score"),
    F("female_agency", "Female Agency Score", "score"),
    F("chemistry", "Chemistry Score", "score"),
    # Not in the SKILL.md column list, but the skill requires a chemistry explanation.
    F("chemistry_explanation", "Chemistry Explanation", "long", width=36, analysis=True),
    F("playful_banter", "Playful Banter", "enum", width=13, choices=YES_NO),
    F("serious_conflict", "Serious Conflict", "enum", width=13, choices=YES_NO),
    F("jealousy", "Jealousy or Possessiveness", "enum", width=14, choices=YES_NO),
    F("deception", "Deception or Mistaken Identity", "enum", width=14, choices=YES_NO),
    F("separation", "Separation or Estrangement", "enum", width=14, choices=YES_NO),
    F("sacrifice", "Sacrifice", "enum", width=13, choices=YES_NO),
    F("team", "Couple Functions as Team", "enum", width=14, choices=YES_NO),
    F("actor_transformation", "Romantic Transformation of Actor’s Character", "long",
      width=36, analysis=True),
    F("female_transformation", "Transformation of Female Character", "long",
      width=36, analysis=True),
    F("relationship_outcome", "Relationship Outcome", "long", width=30),
    F("long_term_union", "Marriage or Long-Term Union", "enum", width=14, choices=YES_NO),
    F("ending", "Happy, Tragic or Unresolved Ending", "enum", width=16, choices=ENDINGS),
    F("romantic_songs", "Major Romantic Songs", "long", width=30),
    F("relationship_analysis", "Relationship Analysis", "long", width=60, analysis=True),
    F("critical_reading", "Contemporary Critical Reading", "long", width=50, analysis=True),
    F("source_ids", "Source IDs", "refs", ref="sources"),
    F("confidence", "Confidence Grade", "enum", width=12, choices=CONFIDENCE),
]

COLLABORATOR_FIELDS = [
    F("person_id", "Person ID", "id"),
    F("name", "Name", width=24),
    F("function", "Profession or Function", width=22),
    F("birth_year", "Birth Year", "year"),
    F("death_year", "Death Year", "year"),
    F("years_active", "Years Active", width=13),
    F("industry", "Industry or Primary Language", width=18),
    F("career_background", "Career Background", "long", width=40),
    F("debut", "Debut or Early Breakthrough", "long", width=30),
    F("major_works", "Major Works", "long", width=36),
    F("recurring_themes", "Recurring Genres or Themes", "long", width=32, analysis=True),
    F("creative_style", "Creative Style", "long", width=36, analysis=True),
    F("major_awards", "Major Awards", "long", width=32),
    F("historical_importance", "Historical Importance", "long", width=36, analysis=True),
    F("relationship_with_actor", "Relationship with Selected Actor", "long", width=36,
      analysis=True),
    F("collaboration_count", "Number of Collaborations", "int", width=14),
    F("collaboration_film_ids", "Collaboration Film IDs", "refs", width=26, ref="films"),
    F("most_important_collaboration", "Most Important Collaboration", width=26),
    F("biography", "Extended Biography", "long", width=90),
    F("source_ids", "Source IDs", "refs", ref="sources"),
    F("confidence", "Confidence Grade", "enum", width=12, choices=CONFIDENCE),
]

CAST_FIELDS = [
    F("cast_id", "Cast ID", "id"),
    F("film_id", "Film ID", "ref", ref="films"),
    F("film", "Film", width=24),
    F("release_year", "Release Year", "year"),
    F("actor", "Actor", width=22),
    F("character", "Character", width=20),
    F("cast_function", "Cast Function", "enum", width=16, choices=CAST_FUNCTIONS),
    F("importance", "Importance Level", "enum", width=12, choices=IMPORTANCE),
    F("relationship_to_protagonist", "Relationship to Protagonist", "long", width=30),
    F("relationship_to_romance", "Relationship to Romantic Plot", "long", width=30),
    F("character_gist", "Character Gist", "long", width=50, analysis=True),
    F("performance_recognition", "Performance Recognition", "long", width=30),
    F("source_ids", "Source IDs", "refs", ref="sources"),
]

AWARD_FIELDS = [
    F("award_id", "Award ID", "id"),
    F("film_id", "Film ID", "ref", ref="films"),
    F("film", "Film", width=24),
    F("ceremony_year", "Ceremony Year", "year"),
    F("organization", "Award Organization", width=26),
    F("category", "Category", width=30),
    F("recipient", "Recipient", width=24),
    F("recipient_type", "Recipient Type", width=16),
    F("result", "Win or Nomination", "enum", width=14, choices=WIN_NOM),
    F("competitive", "Competitive or Honorary", "enum", width=14, choices=COMP_HON),
    F("notes", "Notes", "long", width=36),
    F("source_ids", "Source IDs", "refs", ref="sources"),
    F("confidence", "Confidence Grade", "enum", width=12, choices=CONFIDENCE),
]

SOURCE_FIELDS = [
    F("source_id", "Source ID", "id"),
    F("film_ids", "Film IDs Supported", "refs", ref="films"),
    F("pairing_ids", "Pairing IDs Supported", "refs", ref="romances"),
    F("title", "Source Title", "long", width=40),
    F("publisher", "Publisher", width=22),
    F("author", "Author", width=20),
    F("publication_date", "Publication Date", "date"),
    F("url", "URL", "url"),
    F("access_date", "Access Date", "date"),
    F("claims", "Claims Supported", "long", width=40),
    F("source_type", "Source Type", width=18),
    F("reliability", "Reliability Tier", "enum", width=10, choices=RELIABILITY),
    F("notes", "Notes", "long", width=36),
]

SHEETS = [
    SheetSpec("films", "Films", "films.json", "FILM", "C2", FILM_FIELDS,
              ["title", "release_year", "language", "director", "character", "role_type",
               "role_gist", "romance_presence", "source_ids", "confidence"]),
    SheetSpec("romances", "Romances", "romances.json", "PAIR", "D2", ROMANCE_FIELDS,
              ["film_id", "actress", "relationship_category", "pair_friction",
               "external_opposition", "lover_boy", "romance_centrality",
               "relationship_health", "mutuality", "female_agency", "chemistry",
               "relationship_analysis", "source_ids", "confidence"]),
    SheetSpec("collaborators", "Collaborators", "collaborators.json", "PERSON", "C2",
              COLLABORATOR_FIELDS, ["name", "function", "collaboration_film_ids",
                                    "biography", "source_ids", "confidence"]),
    SheetSpec("cast", "Ensemble Cast", "ensemble_cast.json", "CAST", "D2", CAST_FIELDS,
              ["film_id", "actor", "character", "cast_function", "importance", "source_ids"]),
    SheetSpec("awards", "Awards", "awards.json", "AWARD", "C2", AWARD_FIELDS,
              ["ceremony_year", "organization", "category", "recipient", "result",
               "source_ids", "confidence"]),
    SheetSpec("sources", "Sources", "sources.json", "SOURCE", "B2", SOURCE_FIELDS,
              ["title", "publisher", "reliability", "claims"]),
]
SPEC = {s.key: s for s in SHEETS}
SHEET_ORDER = ["Overview"] + [s.title for s in SHEETS]
UNRELEASED_FILE = "unreleased.json"  # optional; Films schema, IDs UNREL-001, no money
META_FILE = "meta.json"
# Denormalised columns copied from Films when blank (and checked when filled).
FILM_LOOKUPS = {"film": "title", "release_year": "release_year"}

# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------

class Issues:
    def __init__(self) -> None:
        self.errors: list[str] = []
        self.warnings: list[str] = []

    def error(self, where: str, msg: str) -> None:
        self.errors.append(f"{where}: {msg}")

    def warn(self, where: str, msg: str) -> None:
        self.warnings.append(f"{where}: {msg}")


@dataclass
class Dataset:
    meta: dict
    rows: dict[str, list[dict]]
    unreleased: list[dict]


def blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip()) or (isinstance(v, list) and not v)


def is_number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def split_multi(v) -> list[str]:
    if blank(v):
        return []
    parts = v if isinstance(v, list) else str(v).split(";")
    return [str(p).strip() for p in parts if str(p).strip()]


def load_json(path: Path, issues: Issues, expect: type):
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        issues.error(path.name, "file is missing (run `init` to create the research files)")
        return None
    except json.JSONDecodeError as e:
        issues.error(path.name, f"invalid JSON at line {e.lineno}, column {e.colno}: {e.msg}")
        return None
    if not isinstance(data, expect):
        issues.error(path.name, f"top level must be a JSON {'array' if expect is list else 'object'}")
        return None
    return data


def load_rows(path: Path, issues: Issues, required: bool = True) -> list[dict]:
    if not required and not path.exists():
        return []
    data = load_json(path, issues, list) or []
    rows = []
    for i, row in enumerate(data, 1):
        if isinstance(row, dict):
            rows.append(row)
        else:
            issues.error(f"{path.name} item {i}", "each item must be a JSON object")
    return rows


def load_dataset(data_dir: Path, issues: Issues) -> Dataset:
    if not data_dir.is_dir():
        issues.error(str(data_dir), "data directory does not exist")
        return Dataset({}, {s.key: [] for s in SHEETS}, [])
    meta = load_json(data_dir / META_FILE, issues, dict) or {}
    rows = {s.key: load_rows(data_dir / s.file, issues) for s in SHEETS}
    unreleased = load_rows(data_dir / UNRELEASED_FILE, issues, required=False)
    return Dataset(meta, rows, unreleased)

# ---------------------------------------------------------------------------
# Validation (normalises values in place)
# ---------------------------------------------------------------------------

DATE_RE = re.compile(r"^\d{4}(-\d{2}(-\d{2})?)?$")


def suggest(value: str, options) -> str:
    close = difflib.get_close_matches(str(value), list(options), n=1, cutoff=0.6)
    return f" (did you mean {close[0]!r}?)" if close else ""


def check_value(f: Field, v, where: str, issues: Issues, ids: dict[str, set]):
    """Validate one cell and return its normalised value (None when blank)."""
    if blank(v):
        return None
    where = f"{where} {f.key}"
    k = f.kind
    if k in ("refs", "multi"):
        parts = split_multi(v)
        for p in parts:
            if k == "refs" and p not in ids[f.ref]:
                issues.error(where, f"{p!r} does not exist in {SPEC[f.ref].title}")
            if k == "multi" and p not in f.choices:
                issues.error(where, f"{p!r} is not an allowed value{suggest(p, f.choices)}")
        return "; ".join(parts)
    if isinstance(v, list):
        if k in ("text", "long"):
            return "; ".join(split_multi(v))
        issues.error(where, "a list is only allowed in multi-valued columns")
        return None
    if k == "score":
        if not is_number(v) or not 0 <= v <= 10:
            issues.error(where, f"{v!r} must be a number from 0 to 10")
    elif k == "year":
        if not isinstance(v, int) or isinstance(v, bool) or not 1880 <= v <= 2100:
            issues.error(where, f"{v!r} must be a four-digit year as a number")
    elif k == "int":
        if not isinstance(v, int) or isinstance(v, bool) or v < 0:
            issues.error(where, f"{v!r} must be a whole number")
    elif k == "money":
        if is_number(v) and v < 0:
            issues.error(where, "amount cannot be negative")
        elif not is_number(v) and not isinstance(v, str):
            issues.error(where, "use a number, or text for a range or disputed figure")
    elif k == "date":
        s = str(v).strip()
        valid = bool(DATE_RE.match(s))
        if valid and len(s) == 10:
            try:
                dt.date.fromisoformat(s)
            except ValueError:
                valid = False
        if not valid:
            issues.error(where, f"{v!r} must be YYYY-MM-DD (or YYYY-MM / YYYY when only that is known)")
        return s
    elif k == "ref":
        if str(v) not in ids[f.ref]:
            issues.error(where, f"{v!r} does not exist in {SPEC[f.ref].title}")
        return str(v)
    elif k == "enum":
        if v not in f.choices:
            issues.error(where, f"{v!r} is not an allowed value{suggest(v, f.choices)}")
    elif k == "url":
        if not re.match(r"^https?://\S+$", str(v)):
            issues.error(where, f"{v!r} must be a full http(s) URL")
    elif k in ("text", "long", "id") and not isinstance(v, str):
        return str(v)
    return v.strip() if isinstance(v, str) else v


def check_rows(spec: SheetSpec, rows: list[dict], issues: Issues, ids: dict[str, set],
               title: str, prefix: str, recommended: bool = True) -> None:
    keys = {f.key for f in spec.fields}
    id_key = spec.fields[0].key
    pattern = re.compile(rf"^{prefix}-\d{{3,}}$")
    seen: set[str] = set()
    blanks: dict[str, list[str]] = defaultdict(list)
    for i, row in enumerate(rows, 1):
        rid = row.get(id_key)
        label = str(rid) if not blank(rid) else f"row {i}"
        where = f"{title} {label}"
        if blank(rid):
            issues.error(where, f"{id_key} is required")
        elif not pattern.match(str(rid)):
            issues.error(where, f"{id_key} must look like {prefix}-001")
        elif rid in seen:
            issues.error(where, f"duplicate {id_key}")
        seen.add(str(rid))
        for key in row:
            if key not in keys:
                issues.error(where, f"unknown key {key!r}{suggest(key, keys)}")
        for f in spec.fields[1:]:
            row[f.key] = check_value(f, row.get(f.key), where, issues, ids)
        if spec.key == "films":
            derive_film_dates(row, where, issues)  # before blank checks, so derived years count
        for key in spec.recommended if recommended else []:
            if blank(row.get(key)):
                blanks[key].append(label)
    for key, labels in blanks.items():
        shown = ", ".join(labels[:5]) + (f" and {len(labels) - 5} more" if len(labels) > 5 else "")
        issues.warn(title, f"{key} is blank in {len(labels)} row(s): {shown}")


def validate_meta(meta: dict, issues: Issues) -> None:
    if blank(meta.get("actor")):
        issues.error(META_FILE, "actor is required")
    cutoff = meta.get("research_cutoff")
    if not blank(cutoff):
        try:
            dt.date.fromisoformat(str(cutoff))
        except ValueError:
            issues.error(META_FILE, "research_cutoff must be YYYY-MM-DD")
    for i, phase in enumerate(meta.get("career_phases") or [], 1):
        where = f"{META_FILE} career_phases[{i}]"
        if not isinstance(phase, dict) or blank(phase.get("phase")):
            issues.error(where, "each phase needs at least a 'phase' name")
            continue
        for key in ("start_year", "end_year"):
            if not isinstance(phase.get(key), int):
                issues.error(where, f"{key} must be a year as a number")


def validate(ds: Dataset, issues: Issues) -> None:
    validate_meta(ds.meta, issues)
    ids = {s.key: {str(r.get(s.fields[0].key)) for r in ds.rows[s.key]
                   if not blank(r.get(s.fields[0].key))} for s in SHEETS}
    for spec in SHEETS:
        check_rows(spec, ds.rows[spec.key], issues, ids, spec.title, spec.prefix)
    if ds.unreleased:
        check_rows(SPEC["films"], ds.unreleased, issues, ids, "Unreleased", "UNREL", recommended=False)
        for row in ds.unreleased:
            filled = [k for k in MONEY_KEYS if not blank(row.get(k))]
            if filled:
                issues.error(f"Unreleased {row.get('film_id')}",
                             f"commercial figures are not allowed for unreleased films: {', '.join(filled)}")
    check_films(ds, issues)
    check_linked_rows(ds, issues)
    check_collaborators(ds, issues)
    check_sources(ds, issues)

def id_num(value) -> int:
    m = re.search(r"-(\d+)$", str(value or ""))
    return int(m.group(1)) if m else 10**9


def derive_film_dates(row: dict, where: str, issues: Issues) -> None:
    date, year = row.get("release_date"), row.get("release_year")
    if isinstance(date, str) and DATE_RE.match(date):
        date_year = int(date[:4])
        if year is None:
            row["release_year"] = year = date_year
        elif year != date_year:
            issues.error(where, f"release_year {year} does not match release_date {date}")
    if isinstance(year, int):
        decade = f"{year // 10 * 10}s"
        if row.get("decade") is None:
            row["decade"] = decade
        elif row["decade"] != decade:
            issues.warn(where, f"decade {row['decade']!r} replaced with {decade!r}")
            row["decade"] = decade


def film_sort_key(row: dict):
    year = row.get("release_year")
    return (year if isinstance(year, int) else 9999, str(row.get("release_date") or ""),
            id_num(row.get("film_id")))


def check_films(ds: Dataset, issues: Issues) -> None:
    films = ds.rows["films"]
    seen: dict[tuple, str] = {}
    for row in films:
        where = f"Films {row.get('film_id')}"
        if row.get("title") and isinstance(row.get("release_year"), int):
            key = (row["title"].casefold(), row["release_year"])
            if key in seen:
                issues.error(where, f"same title and year as {seen[key]}; each film needs exactly one row")
            seen[key] = row.get("film_id")
        if any(is_number(row.get(k)) for k in MONEY_KEYS) and blank(row.get("currency")):
            issues.warn(where, "commercial figures have no Currency")
        nett, dom, world = (row.get(k) for k in ("domestic_nett", "domestic_gross", "worldwide_gross"))
        if is_number(dom) and is_number(world) and world < dom:
            issues.warn(where, "worldwide gross is lower than domestic gross; check nett/gross labels and units")
        if is_number(nett) and is_number(dom) and dom < nett:
            issues.warn(where, "domestic gross is lower than domestic nett; check the labels")

    in_id_order = [r.get("film_id") for r in sorted(films, key=lambda r: id_num(r.get("film_id")))]
    in_date_order = [r.get("film_id") for r in sorted(films, key=film_sort_key)]
    if in_id_order != in_date_order:
        issues.warn("Films", "Film IDs are not in release order (the workbook sorts Films by release date)")

    pairings = Counter(r.get("film_id") for r in ds.rows["romances"])
    for row in films:
        where, presence = f"Films {row.get('film_id')}", row.get("romance_presence")
        count = pairings[row.get("film_id")]
        if presence == "None" and count:
            issues.warn(where, f"romance_presence is None but {count} Romances row(s) exist")
        if presence in ("Central", "Major") and not count:
            issues.warn(where, f"romance_presence is {presence} but there is no Romances row")


def check_linked_rows(ds: Dataset, issues: Issues) -> None:
    films = {r["film_id"]: r for r in ds.rows["films"] if r.get("film_id")}
    for key in ("romances", "cast", "awards"):
        spec = SPEC[key]
        for row in ds.rows[key]:
            film = films.get(row.get("film_id"))
            if not film:
                continue
            where = f"{spec.title} {row.get(spec.fields[0].key)}"
            lookups = dict(FILM_LOOKUPS, **({"character": "character"} if key == "romances" else {}))
            for col, film_col in lookups.items():
                own, master = row.get(col), film.get(film_col)
                if master is None:
                    continue
                if own is None:
                    row[col] = master
                elif col != "character" and str(own).casefold() != str(master).casefold():
                    issues.error(where, f"{col} {own!r} does not match Films ({master!r}); "
                                        "leave it blank to copy it from Films")
            if key == "romances":
                own, master = row.get("romance_centrality"), film.get("romance_centrality")
                if is_number(own) and is_number(master) and own > master:
                    issues.warn(where, f"pairing romance centrality {own} exceeds the film's {master}")


def check_collaborators(ds: Dataset, issues: Issues) -> None:
    for row in ds.rows["collaborators"]:
        n = len(split_multi(row.get("collaboration_film_ids")))
        if row.get("collaboration_count") is None:
            row["collaboration_count"] = n or None
        elif n and row["collaboration_count"] != n:
            issues.warn(f"Collaborators {row.get('person_id')}",
                        f"collaboration_count is {row['collaboration_count']} but {n} Film IDs are listed")


def check_sources(ds: Dataset, issues: Issues) -> None:
    cited: set[str] = set()
    for spec in SHEETS:
        if spec.key != "sources":
            for row in ds.rows[spec.key]:
                cited.update(split_multi(row.get("source_ids")))
    for row in ds.unreleased:
        cited.update(split_multi(row.get("source_ids")))
    unused = [r.get("source_id") for r in ds.rows["sources"] if r.get("source_id") not in cited]
    if unused:
        issues.warn("Sources", f"{len(unused)} source(s) are not cited by any row: {', '.join(unused[:10])}")

# ---------------------------------------------------------------------------
# Overview metrics
# ---------------------------------------------------------------------------
# Each metric carries the cell content (a value or an Excel formula) and the
# value computed in Python. openpyxl cannot calculate formulas, so the build
# report prints the Python value for reconciliation.

@dataclass
class Metric:
    label: str
    cell: object
    expected: object
    method: str  # "Formula" | "Build script" | "Research file"
    note: str = ""


def xl_round(x: float, places: int = 1) -> float:
    return float(Decimal(str(x)).quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP))


def fmt_num(v) -> str:
    return f"{v:,.2f}".rstrip("0").rstrip(".") if is_number(v) else str(v)


def rng(sheet_key: str, field_key: str, n: int) -> str:
    spec = SPEC[sheet_key]
    col = spec.col(field_key)
    title = f"'{spec.title}'" if " " in spec.title else spec.title
    return f"{title}!${col}$2:${col}${max(n, 1) + 1}"


def xl_text(s: str) -> str:
    return '"' + str(s).replace('"', '""') + '"'


def count_formula(r: str, values: list[str]) -> str:
    return "=" + "+".join(f"COUNTIF({r},{xl_text(v)})" for v in values)


def tally_formula(r: str, values: list[str]) -> str:
    parts = [f'{xl_text(v + " ")}&COUNTIF({r},{xl_text(v)})' for v in values]
    return "=" + '&" · "&'.join(parts)


def mean(values) -> float | str:
    nums = [v for v in values if is_number(v)]
    return xl_round(sum(nums) / len(nums)) if nums else "n/a"


def per_film_names(rows: list[dict], key: str) -> list[str]:
    names: list[str] = []
    for row in rows:
        names.extend(dict.fromkeys(split_multi(row.get(key))))  # once per film
    return names


def most_frequent(names: list[str]) -> str:
    counts = Counter(names)
    if not counts:
        return "n/a"
    top = max(counts.values())
    leaders = sorted(n for n, c in counts.items() if c == top)
    each = " each" if len(leaders) > 1 else ""
    return f"{'; '.join(leaders)} ({top} film{'s' if top != 1 else ''}{each})"


def costar_film_counts(romances: list[dict]) -> Counter:
    per_film: dict[str, set] = defaultdict(set)
    for r in romances:
        if r.get("actress"):
            per_film[r.get("film_id")].add(r["actress"])
    return Counter(a for names in per_film.values() for a in names)


def highest_grossing(films: list[dict]) -> tuple[str, str]:
    numeric = [f for f in films if is_number(f.get("worldwide_gross")) and f.get("currency")]
    if not numeric:
        return "n/a", "No numeric worldwide gross with a currency recorded"
    currencies = Counter(f["currency"] for f in numeric)
    currency = currencies.most_common(1)[0][0]
    best = max((f for f in numeric if f["currency"] == currency), key=lambda f: f["worldwide_gross"])
    note = f"Largest nominal worldwide gross among {currencies[currency]} film(s) reported in {currency}; not inflation-adjusted"
    if len(numeric) > currencies[currency]:
        note += f"; {len(numeric) - currencies[currency]} film(s) in other currencies excluded"
    ranged = sum(1 for f in films if isinstance(f.get("worldwide_gross"), str))
    if ranged:
        note += f"; {ranged} film(s) with ranged or disputed figures excluded"
    value = f"{best.get('title')} ({best.get('release_year')}): {fmt_num(best['worldwide_gross'])} {currency}"
    return value, note


def build_metrics(ds: Dataset) -> list[Metric]:
    films, rom = ds.rows["films"], ds.rows["romances"]
    awards, sources = ds.rows["awards"], ds.rows["sources"]
    nf, nr, na, ns = len(films), len(rom), len(awards), len(sources)
    actor = ds.meta.get("actor") or ""
    cutoff = ds.meta.get("research_cutoff")
    try:
        cutoff = dt.date.fromisoformat(str(cutoff)) if cutoff else "n/a"
    except ValueError:
        pass

    years_r = rng("films", "release_year", nf)
    years = [f["release_year"] for f in films if isinstance(f.get("release_year"), int)]
    metrics = [
        Metric("Actor", actor, actor, "Research file"),
        Metric("Career period", f'=IF(COUNT({years_r})=0,"n/a",MIN({years_r})&"–"&MAX({years_r}))',
               f"{min(years)}–{max(years)}" if years else "n/a", "Formula",
               "First to latest theatrical release year in Films"),
        Metric("Research cutoff date", cutoff, cutoff, "Research file"),
        Metric("Total released acting films", f"=COUNTA({rng('films', 'film_id', nf)})", nf, "Formula"),
    ]
    role_r = rng("films", "role_type", nf)
    for label, values in ROLE_GROUPS:
        metrics.append(Metric(label, count_formula(role_r, values),
                              sum(f.get("role_type") in values for f in films), "Formula",
                              "Role Type: " + ", ".join(values)))
    presence_r = rng("films", "romance_presence", nf)
    for label, values in ROMANCE_GROUPS:
        metrics.append(Metric(label, count_formula(presence_r, values),
                              sum(f.get("romance_presence") in values for f in films), "Formula",
                              "Romance Presence: " + ", ".join(values)))

    banner_key = "production_banner" if any(f.get("production_banner") for f in films) else "producers"
    costars = costar_film_counts(rom)
    metrics += [
        Metric("Distinct principal female co-stars", len(set(per_film_names(films, "female_leads"))),
               len(set(per_film_names(films, "female_leads"))), "Build script",
               "Distinct names in Principal Female Lead or Leads"),
        Metric("Romantic pairings recorded", f"=COUNTA({rng('romances', 'pairing_id', nr)})", nr, "Formula"),
        Metric("Most frequent romantic co-star", most_frequent(list(costars.elements())),
               most_frequent(list(costars.elements())), "Build script", "Films with a Romances row for her"),
        Metric("Most frequent director", most_frequent(per_film_names(films, "director")),
               most_frequent(per_film_names(films, "director")), "Build script"),
        Metric("Most frequent producer or production banner",
               most_frequent(per_film_names(films, banner_key)),
               most_frequent(per_film_names(films, banner_key)), "Build script",
               "From Production Banner" if banner_key == "production_banner" else "From Producer or Producers"),
        Metric("Most frequent music director", most_frequent(per_film_names(films, "song_composer")),
               most_frequent(per_film_names(films, "song_composer")), "Build script", "From Song Composer"),
    ]
    for label, sheet, key, rows, n, note in [
        ("Average pair-friction score", "romances", "pair_friction", rom, nr, "Per pairing"),
        ("Average external-opposition score", "romances", "external_opposition", rom, nr, "Per pairing"),
        ("Average lover-boy score", "romances", "lover_boy", rom, nr, "Per pairing"),
        ("Average romance-centrality score", "films", "romance_centrality", films, nf, "Per film"),
    ]:
        metrics.append(Metric(label, f'=IFERROR(ROUND(AVERAGE({rng(sheet, key, n)}),1),"n/a")',
                              mean(r.get(key) for r in rows), "Formula", note))

    top_film, top_note = highest_grossing(films)
    result_r, recipient_r = rng("awards", "result", na), rng("awards", "recipient", na)
    rtype_r = rng("awards", "recipient_type", na)
    actor_pattern = xl_text(f"*{actor}*")
    by_actor = [a for a in awards if actor and actor.casefold() in str(a.get("recipient") or "").casefold()]
    metrics += [
        Metric("Highest-grossing film", top_film, top_film, "Build script", top_note),
        Metric("Award wins received by the actor", f"=COUNTIFS({result_r},\"Win\",{recipient_r},{actor_pattern})",
               sum(a.get("result") == "Win" for a in by_actor), "Formula",
               "Awards rows whose Recipient includes the actor, in any capacity: acting, directing, producing, honours"),
        Metric("Acting award wins",
               f"=COUNTIFS({result_r},\"Win\",{recipient_r},{actor_pattern},{rtype_r},\"Actor\")",
               sum(a.get("result") == "Win" and str(a.get("recipient_type") or "").casefold() == "actor"
                   for a in by_actor), "Formula", "As above, limited to Recipient Type 'Actor'"),
        Metric("Award nominations received by the actor",
               f"=COUNTIFS({result_r},\"Nomination\",{recipient_r},{actor_pattern})",
               sum(a.get("result") == "Nomination" for a in by_actor), "Formula", "Nominations that did not win"),
        Metric("Award wins for his films (all recipients)", f"=COUNTIF({result_r},\"Win\")",
               sum(a.get("result") == "Win" for a in awards), "Formula"),
        Metric("Source-confidence summary (films)", tally_formula(rng("films", "confidence", nf), CONFIDENCE),
               " · ".join(f"{g} {sum(f.get('confidence') == g for f in films)}" for g in CONFIDENCE),
               "Formula", "Confidence Grade in Films"),
        Metric("Sources by reliability tier", tally_formula(rng("sources", "reliability", ns), RELIABILITY),
               " · ".join(f"{t} {sum(s.get('reliability') == t for s in sources)}" for t in RELIABILITY),
               "Formula", "A official · B reputable press · C reference databases · D weak"),
    ]
    return metrics


def phase_rows(ds: Dataset) -> list[dict]:
    films, rom = ds.rows["films"], ds.rows["romances"]
    years_r = rng("films", "release_year", len(films))
    rom_years_r, lover_r = rng("romances", "release_year", len(rom)), rng("romances", "lover_boy", len(rom))
    rows = []
    for p in ds.meta.get("career_phases") or []:
        start, end = p.get("start_year"), p.get("end_year")
        if not isinstance(start, int) or not isinstance(end, int):
            continue
        lo, hi = xl_text(f">={start}"), xl_text(f"<={end}")
        rows.append({
            "phase": p.get("phase"), "years": f"{start}–{end}", "summary": p.get("summary"),
            "films": f"=COUNTIFS({years_r},{lo},{years_r},{hi})",
            "films_expected": sum(start <= (f.get("release_year") or 0) <= end for f in films),
            "lover": f'=IFERROR(ROUND(AVERAGEIFS({lover_r},{rom_years_r},{lo},{rom_years_r},{hi}),1),"n/a")',
            "lover_expected": mean(r.get("lover_boy") for r in rom
                                   if start <= (r.get("release_year") or 0) <= end),
        })
    return rows


def decade_rows(films: list[dict]) -> list[tuple[str, int, object]]:
    by_decade: dict[str, list] = defaultdict(list)
    for f in films:
        if f.get("decade"):
            by_decade[f["decade"]].append(f.get("romance_centrality"))
    return [(d, len(v), mean(v)) for d, v in sorted(by_decade.items())]

# ---------------------------------------------------------------------------
# Data-sheet rendering
# ---------------------------------------------------------------------------

HEADER_FILL = {"fact": "1F3A5F", "score": "8A5A00", "analysis": "3B5E3F"}
ANALYSIS_FILL = "F3F7F0"
SCORE_SCALE = ("FFFFFF", "FCE7B2", "E07B39")  # 0, 5, 10
THIN = Side(style="thin", color="BFC5CC")
LINK_BLUE = "0563C1"
LINE_HEIGHT = 15.0
MAX_ROW_HEIGHT = 409.0
EXTRA_VALIDATION_ROWS = 200
LIST_COLUMN_START = 40  # hidden Overview columns holding choice lists too long to inline
NUMBER_FORMATS = {"score": "General", "year": "0", "int": "0", "num": "General",
                  "money": "#,##0.00", "date": "yyyy-mm-dd"}
CENTERED = {"id", "ref", "score", "year", "int", "num"}


def text_lines(text, width: float, bold: bool = False) -> int:
    """Estimate wrapped line count; errs high so cells are not clipped."""
    per_line = max(4, int(width * (1.05 if bold else 1.2)))
    return sum(max(1, math.ceil(len(part) / per_line)) for part in str(text).split("\n"))


class ChoiceLists:
    """Inline list validations where Excel allows it (<=255 chars), else hidden Overview columns."""

    def __init__(self, overview) -> None:
        self.ws = overview
        self.refs: dict[tuple, str] = {}
        self.next_col = LIST_COLUMN_START

    def formula(self, name: str, choices: list[str]) -> str:
        inline = '"' + ",".join(choices) + '"'
        if len(inline) <= 250 and not any("," in c for c in choices):
            return inline
        key = tuple(choices)
        if key not in self.refs:
            col = get_column_letter(self.next_col)
            self.ws.cell(1, self.next_col, name)
            for i, choice in enumerate(choices, 2):
                self.ws.cell(i, self.next_col, choice)
            self.ws.column_dimensions[col].hidden = True
            self.refs[key] = f"Overview!${col}$2:${col}${len(choices) + 1}"
            self.next_col += 1
        return self.refs[key]


def write_header(ws, row: int, fields: list[Field]) -> None:
    lines = 1
    for c, f in enumerate(fields, 1):
        cell = ws.cell(row, c, f.header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL[f.category])
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = Border(bottom=THIN)
        lines = max(lines, text_lines(f.header, f.col_width, bold=True))
    ws.row_dimensions[row].height = min(MAX_ROW_HEIGHT, lines * LINE_HEIGHT + 6)


def write_rows(ws, fields: list[Field], rows: list[dict], start_row: int) -> None:
    analysis_fill = PatternFill("solid", fgColor=ANALYSIS_FILL)
    for r, row in enumerate(rows, start_row):
        lines = 1
        for c, f in enumerate(fields, 1):
            value = row.get(f.key)
            if f.kind == "date" and isinstance(value, str) and len(value) == 10:
                value = dt.date.fromisoformat(value)
            cell = ws.cell(r, c)
            cell.value = value
            if isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"  # text, not a formula
            numeric = value is not None and not isinstance(value, str)
            if numeric and f.kind in NUMBER_FORMATS:
                cell.number_format = NUMBER_FORMATS[f.kind]
            horizontal = "center" if f.kind in CENTERED else ("right" if numeric else "left")
            cell.alignment = Alignment(wrap_text=not numeric, vertical="top", horizontal=horizontal)
            if f.kind == "id":
                cell.font = Font(bold=True)
            elif f.kind == "url" and value:
                cell.hyperlink = value
                cell.font = Font(color=LINK_BLUE, underline="single")
            if f.category == "analysis":
                cell.fill = analysis_fill
            if value is not None and not numeric:
                lines = max(lines, text_lines(value, f.col_width))
        ws.row_dimensions[r].height = min(MAX_ROW_HEIGHT, lines * LINE_HEIGHT + 4)


def add_validation(ws, fields: list[Field], first_row: int, last_row: int, lists: ChoiceLists) -> None:
    for c, f in enumerate(fields, 1):
        col = get_column_letter(c)
        ref = f"{col}{first_row}:{col}{last_row}"
        if f.kind == "enum":
            dv = DataValidation(type="list", formula1=lists.formula(f.header, f.choices),
                                allow_blank=True, showErrorMessage=True, errorTitle=f.header,
                                error="Choose a value from the list.")
        elif f.kind == "score":
            dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="10",
                                allow_blank=True, showErrorMessage=True, errorTitle=f.header,
                                error="Enter a score from 0 to 10.")
            start, mid, end = SCORE_SCALE
            ws.conditional_formatting.add(ref, ColorScaleRule(
                start_type="num", start_value=0, start_color=start,
                mid_type="num", mid_value=5, mid_color=mid,
                end_type="num", end_value=10, end_color=end))
        else:
            continue
        dv.add(ref)
        ws.add_data_validation(dv)


def write_data_sheet(ws, spec: SheetSpec, rows: list[dict], lists: ChoiceLists,
                     unreleased: list[dict] | None = None) -> None:
    for c, f in enumerate(spec.fields, 1):
        ws.column_dimensions[get_column_letter(c)].width = f.col_width
    write_header(ws, 1, spec.fields)
    write_rows(ws, spec.fields, rows, 2)
    last_row = max(len(rows), 1) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(spec.fields))}{last_row}"
    ws.freeze_panes = spec.freeze
    extra = 0 if unreleased else EXTRA_VALIDATION_ROWS
    add_validation(ws, spec.fields, 2, last_row + extra, lists)
    if unreleased:
        label_row = last_row + 2
        ws.cell(label_row, 1, "Unreleased or abandoned projects (excluded from all totals and charts)"
                ).font = Font(bold=True, italic=True, size=12)
        write_header(ws, label_row + 1, spec.fields)
        write_rows(ws, spec.fields, unreleased, label_row + 2)
        add_validation(ws, spec.fields, label_row + 2, label_row + 1 + len(unreleased), lists)


def sorted_rows(ds: Dataset) -> dict[str, list[dict]]:
    films = sorted(ds.rows["films"], key=film_sort_key)
    order = {f.get("film_id"): i for i, f in enumerate(films)}
    by_film = lambda id_key: (lambda r: (order.get(r.get("film_id"), len(order)), id_num(r.get(id_key))))
    return {
        "films": films,
        "romances": sorted(ds.rows["romances"], key=by_film("pairing_id")),
        "collaborators": sorted(ds.rows["collaborators"], key=lambda r: id_num(r.get("person_id"))),
        "cast": sorted(ds.rows["cast"], key=by_film("cast_id")),
        "awards": sorted(ds.rows["awards"],
                         key=lambda r: (r.get("ceremony_year") or 9999, id_num(r.get("award_id")))),
        "sources": sorted(ds.rows["sources"], key=lambda r: id_num(r.get("source_id"))),
    }

# ---------------------------------------------------------------------------
# Overview rendering
# ---------------------------------------------------------------------------

OVERVIEW_WIDTHS = {"A": 40, "B": 38, "C": 14, "D": 52, "E": 14, "F": 3}
TITLE_FONT = Font(bold=True, size=16, color="1F3A5F")
SECTION_FONT = Font(bold=True, size=12, color="1F3A5F")
CHART_COLUMN = "G"
CHART_ROW_STEP = 17


def col_idx(sheet_key: str, field_key: str) -> int:
    return next(i for i, f in enumerate(SPEC[sheet_key].fields, 1) if f.key == field_key)


def overview_header(ws, row: int, labels: list[str]) -> int:
    for c, label in enumerate(labels, 1):
        cell = ws.cell(row, c, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL["fact"])
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    return row + 1


def overview_row(ws, row: int, values: list, formats: dict[int, str] | None = None) -> int:
    """formats maps a 1-based column number to a number format for non-text cells."""
    lines = 1
    for c, value in enumerate(values, 1):
        cell = ws.cell(row, c, value)
        is_text = isinstance(value, str) and not value.startswith("=")
        cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        if isinstance(value, dt.date):
            cell.number_format = "yyyy-mm-dd"
        elif formats and c in formats and not is_text:
            cell.number_format = formats[c]
        if is_text:
            width = OVERVIEW_WIDTHS[get_column_letter(c)]
            lines = max(lines, text_lines(value, width))
    ws.row_dimensions[row].height = min(MAX_ROW_HEIGHT, lines * LINE_HEIGHT + 4)
    return row + 1


def section(ws, row: int, title: str) -> int:
    ws.cell(row, 1, title).font = SECTION_FONT
    return row + 1


def scatter_chart(src, x_col: int, y_col: int, n: int, title: str, x_title: str, y_title: str,
                  x_range: tuple[float, float]) -> ScatterChart:
    chart = ScatterChart()
    chart.title, chart.style, chart.height, chart.width = title, 13, 7.5, 16
    chart.x_axis.title, chart.y_axis.title = x_title, y_title
    series = Series(Reference(src, min_col=y_col, min_row=2, max_row=n + 1),
                    Reference(src, min_col=x_col, min_row=2, max_row=n + 1), title=y_title)
    series.marker.symbol, series.marker.size = "circle", 7
    series.graphicalProperties.line.noFill = True
    chart.series.append(series)
    chart.legend = None
    chart.x_axis.scaling.min, chart.x_axis.scaling.max = x_range
    chart.y_axis.scaling.min, chart.y_axis.scaling.max = 0, 10
    chart.x_axis.delete = chart.y_axis.delete = False  # openpyxl 3.1 hides axes otherwise
    return chart


def bar_chart(ws, header_row: int, last_row: int, value_col: int, title: str, horizontal: bool,
              y_max: float | None = None) -> BarChart:
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.title, chart.style, chart.height, chart.width = title, 10, 7.5, 16
    chart.add_data(Reference(ws, min_col=value_col, min_row=header_row, max_row=last_row),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=header_row + 1, max_row=last_row))
    chart.legend = None
    if horizontal:
        chart.x_axis.scaling.orientation = "maxMin"  # first row at the top
    if y_max is not None:
        chart.y_axis.scaling.min, chart.y_axis.scaling.max = 0, y_max
    chart.x_axis.delete = chart.y_axis.delete = False
    return chart


def render_overview(wb, ds: Dataset, rows: dict[str, list[dict]], metrics: list[Metric]) -> None:
    ws = wb["Overview"]
    for col, width in OVERVIEW_WIDTHS.items():
        ws.column_dimensions[col].width = width
    actor = ds.meta.get("actor") or "Actor"
    ws["A1"] = f"{actor}: Filmography Overview"
    ws["A1"].font = TITLE_FONT
    subtitle = (f"Generated {dt.date.today().isoformat()} from the research files. "
                f"Research cutoff: {ds.meta.get('research_cutoff') or 'not stated'}.")
    if ds.meta.get("scope_note"):
        subtitle += f" {ds.meta['scope_note']}"
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, color="555555")

    row = section(ws, 4, "Career summary")
    row = overview_header(ws, row, ["Metric", "Value", "Method", "Notes"])
    for m in metrics:
        row = overview_row(ws, row, [m.label, m.cell, m.method, m.note],
                           {2: "0.0"} if m.label.startswith("Average") else None)

    row = section(ws, row + 1, "Career phases")
    row = overview_header(ws, row, ["Phase", "Years", "Films", "Summary", "Avg lover-boy score"])
    phases = phase_rows(ds)
    for p in phases:
        row = overview_row(ws, row, [p["phase"], p["years"], p["films"], p["summary"], p["lover"]], {3: "0", 5: "0.0"})
    if not phases:
        row = overview_row(ws, row, ["No career phases recorded in meta.json"])

    costars = costar_film_counts(rows["romances"]).most_common(10)
    row = section(ws, row + 1, "Romantic co-stars (top 10)")
    costar_header = row
    row = overview_header(ws, row, ["Actress", "Films with a pairing"])
    for name, count in costars:
        row = overview_row(ws, row, [name, count])
    costar_last = row - 1

    decades = decade_rows(rows["films"])
    row = section(ws, row + 1, "Romance centrality by decade")
    decade_header = row
    row = overview_header(ws, row, ["Decade", "Films", "Avg romance centrality"])
    for decade, count, avg in decades:
        row = overview_row(ws, row, [decade, count, avg], {2: "0", 3: "0.0"})
    decade_last = row - 1

    row = section(ws, row + 1, "Column colour key")
    for category, text in [("fact", "Fact: credits, dates, figures and descriptive fields"),
                           ("score", "Score (0–10): cells shaded white (0) to orange (10)"),
                           ("analysis", "Interpretive analysis: header green, cells tinted green")]:
        cell = ws.cell(row, 1, category.capitalize())
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL[category])
        ws.cell(row, 2, text)
        row += 1

    charts = []
    rom_ws, n_rom = wb["Romances"], len(rows["romances"])
    rom_years = [r["release_year"] for r in rows["romances"] if isinstance(r.get("release_year"), int)]
    if n_rom and rom_years:
        charts.append(scatter_chart(rom_ws, col_idx("romances", "release_year"), col_idx("romances", "lover_boy"),
                                    n_rom, "Lover-boy score by release year (per pairing)", "Release year",
                                    "Lover-boy score", (min(rom_years) - 1, max(rom_years) + 1)))
    if costars:
        charts.append(bar_chart(ws, costar_header, costar_last, 2, "Romantic co-star frequency", True))
    if n_rom:
        charts.append(scatter_chart(rom_ws, col_idx("romances", "pair_friction"),
                                    col_idx("romances", "external_opposition"), n_rom,
                                    "Pair friction vs external opposition", "Pair friction",
                                    "External opposition", (0, 10)))
    if decades:
        charts.append(bar_chart(ws, decade_header, decade_last, 3, "Average romance centrality by decade",
                                False, y_max=10))
    for i, chart in enumerate(charts):
        ws.add_chart(chart, f"{CHART_COLUMN}{4 + i * CHART_ROW_STEP}")

# ---------------------------------------------------------------------------
# Build, structural self-check and report
# ---------------------------------------------------------------------------

def build_workbook(ds: Dataset, out: Path) -> list[Metric]:
    rows = sorted_rows(ds)
    wb = openpyxl.Workbook()
    wb.active.title = "Overview"
    for spec in SHEETS:
        wb.create_sheet(spec.title)
    lists = ChoiceLists(wb["Overview"])
    for spec in SHEETS:
        unreleased = sorted(ds.unreleased, key=film_sort_key) if spec.key == "films" else None
        write_data_sheet(wb[spec.title], spec, rows[spec.key], lists, unreleased)
    metrics = build_metrics(ds)
    render_overview(wb, ds, rows, metrics)
    wb.save(out)
    return metrics


def self_check(ds: Dataset, out: Path) -> list[str]:
    """Re-open the saved file and confirm the workbook contract."""
    problems = []
    wb = openpyxl.load_workbook(out)
    if wb.sheetnames != SHEET_ORDER:
        problems.append(f"sheet order is {wb.sheetnames}, expected {SHEET_ORDER}")
        return problems
    for spec in SHEETS:
        ws, n = wb[spec.title], len(ds.rows[spec.key])
        if [c.value for c in ws[1]][:len(spec.fields)] != [f.header for f in spec.fields]:
            problems.append(f"{spec.title}: header row does not match the schema")
        written = sorted(str(ws.cell(r, 1).value) for r in range(2, n + 2))
        expected = sorted(str(r.get(spec.fields[0].key)) for r in ds.rows[spec.key])
        if written != expected:
            problems.append(f"{spec.title}: IDs in rows 2–{n + 1} do not match the research file")
        if not ws.auto_filter.ref:
            problems.append(f"{spec.title}: no filter")
        if not ws.freeze_panes:
            problems.append(f"{spec.title}: header row is not frozen")
        if ws.merged_cells.ranges:
            problems.append(f"{spec.title}: contains merged cells")
    films_ws, year_col = wb["Films"], col_idx("films", "release_year")
    years = [films_ws.cell(r, year_col).value for r in range(2, len(ds.rows["films"]) + 2)]
    known = [y for y in years if isinstance(y, int)]
    if known != sorted(known):
        problems.append("Films: rows are not in chronological order")
    for row in wb["Overview"].iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                for quoted, bare in re.findall(r"(?:'([^']+)'|([A-Za-z]+))!\$", cell.value):
                    if (quoted or bare) not in wb.sheetnames:
                        problems.append(f"Overview {cell.coordinate}: refers to unknown sheet {quoted or bare!r}")
    return problems


def check_cutoff(ds: Dataset, issues: Issues) -> None:
    try:
        cutoff = dt.date.fromisoformat(str(ds.meta.get("research_cutoff")))
    except ValueError:
        return
    for row in ds.rows["films"]:
        date = row.get("release_date")
        if isinstance(date, str) and len(date) == 10 and DATE_RE.match(date) and date > cutoff.isoformat():
            issues.warn(f"Films {row.get('film_id')}", f"released {date}, after the research cutoff {cutoff}")


def run_validation(data_dir: str) -> tuple[Dataset, Issues]:
    issues = Issues()
    ds = load_dataset(Path(data_dir), issues)
    if not issues.errors:  # skip cascading reference errors when a file failed to load
        validate(ds, issues)
        check_cutoff(ds, issues)
    return ds, issues


def print_issues(issues: Issues) -> None:
    for e in issues.errors:
        print(f"ERROR    {e}")
    for w in issues.warnings:
        print(f"WARNING  {w}")
    print(f"{len(issues.errors)} error(s), {len(issues.warnings)} warning(s)")


def print_report(ds: Dataset, out: Path, metrics: list[Metric]) -> None:
    counts = ", ".join(f"{s.title} {len(ds.rows[s.key])}" for s in SHEETS)
    if ds.unreleased:
        counts += f", Unreleased {len(ds.unreleased)}"
    print(f"\nWrote {out}\nRows: {counts}")
    print("\nOverview values (formula cells are calculated when the workbook is opened):")
    width = max(len(m.label) for m in metrics)
    for m in metrics:
        value = m.expected.isoformat() if isinstance(m.expected, dt.date) else m.expected
        print(f"  {m.label:<{width}}  {value}")

# ---------------------------------------------------------------------------
# Commands
# ---------------------------------------------------------------------------

KIND_HELP = ("id: sheet ID | ref: one ID | refs: IDs separated by '; ' or a JSON array | "
             "text / long: text | score: number 0–10 | date: YYYY-MM-DD (or YYYY-MM, YYYY) | "
             "year: number | int: whole number | num: number or text | "
             "money: number in the Currency unit, or text for a range | enum: one listed value | "
             "multi: listed values separated by '; ' | url: http(s) link")


def cmd_init(args) -> int:
    data_dir = Path(args.data_dir)
    if args.cutoff:
        try:
            dt.date.fromisoformat(args.cutoff)
        except ValueError:
            print("--cutoff must be YYYY-MM-DD")
            return 2
    data_dir.mkdir(parents=True, exist_ok=True)
    files: dict[str, object] = {META_FILE: {
        "actor": args.actor,
        "research_cutoff": args.cutoff or dt.date.today().isoformat(),
        "scope_note": "Released feature-film acting appearances through the research cutoff.",
        "career_phases": [],
    }}
    files.update({s.file: [] for s in SHEETS})
    for name, content in files.items():
        path = data_dir / name
        if path.exists():
            print(f"kept     {path} (already exists)")
            continue
        path.write_text(json.dumps(content, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        print(f"created  {path}")
    return 0


def cmd_schema(args) -> int:
    aliases = {}
    for s in SHEETS:
        for alias in (s.key, s.title, Path(s.file).stem):
            aliases[alias.lower().replace(" ", "_")] = s
    specs = SHEETS
    if args.sheet:
        spec = aliases.get(args.sheet.lower().replace(" ", "_"))
        if not spec:
            print(f"Unknown sheet {args.sheet!r}. Choose from: {', '.join(s.key for s in SHEETS)}")
            return 2
        specs = [spec]
    else:
        print(f"{META_FILE}: actor, research_cutoff (YYYY-MM-DD), scope_note, "
              "career_phases [{phase, start_year, end_year, summary}]")
        print(f"{UNRELEASED_FILE} (optional): same keys as films, IDs UNREL-001, no commercial figures")
    print(f"Kinds: {KIND_HELP}")
    for spec in specs:
        print(f"\n{spec.title} ({spec.file}, IDs {spec.prefix}-001)")
        for f in spec.fields:
            detail = f.kind
            if f.ref:
                detail += f" → {SPEC[f.ref].title}"
            if f.choices:
                detail += ": " + " | ".join(f.choices)
            print(f"  {f.key:<32} {f.header:<46} {detail}")
    return 0


def cmd_validate(args) -> int:
    _, issues = run_validation(args.data_dir)
    print_issues(issues)
    return 1 if issues.errors else 0


def cmd_build(args) -> int:
    out = Path(args.output)
    if out.suffix.lower() != ".xlsx":
        print("Output file must end in .xlsx")
        return 2
    ds, issues = run_validation(args.data_dir)
    print_issues(issues)
    if issues.errors:
        print("Build stopped: fix the errors above.")
        return 1
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        metrics = build_workbook(ds, out)
    except PermissionError:
        print(f"Cannot write {out}. Close it if it is open in a spreadsheet application.")
        return 1
    print_report(ds, out, metrics)
    problems = self_check(ds, out)
    if problems:
        print("\nStructural check failed:")
        for p in problems:
            print(f"  {p}")
        return 1
    print("\nStructural check passed.")
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build the Bollywood actor filmography workbook.")
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("init", help="create empty research files for an actor")
    p.add_argument("data_dir")
    p.add_argument("--actor", required=True)
    p.add_argument("--cutoff", help="research cutoff date, YYYY-MM-DD (default: today)")
    p.set_defaults(func=cmd_init)
    p = sub.add_parser("schema", help="list JSON keys, kinds and allowed values")
    p.add_argument("sheet", nargs="?", help="films, romances, collaborators, cast, awards or sources")
    p.set_defaults(func=cmd_schema)
    p = sub.add_parser("validate", help="check the research files without building")
    p.add_argument("data_dir")
    p.set_defaults(func=cmd_validate)
    p = sub.add_parser("build", help="validate, then write the workbook")
    p.add_argument("data_dir")
    p.add_argument("-o", "--output", required=True)
    p.set_defaults(func=cmd_build)
    args = parser.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
